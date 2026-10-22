#!/usr/bin/env python3
"""
Consolida todas as métricas estatísticas geradas nos experimentos do Colorímetro
em uma única pasta de trabalho do Excel (.xlsx), com aba de índice e mapeamento
exato de qual script gerou cada tabela.
"""

import os
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = "/media/workstation_2x_3080/ext_ssd1/Colorímetro"
RD_DIR = os.path.join(BASE_DIR, "resultados_estatisticos")
OUTPUT_XLSX = os.path.join(BASE_DIR, "Consolidado_Metricas_Estatisticas_Colorimetro.xlsx")

# Definição das tabelas e seus metadados de rastreabilidade
TABLES_METADATA = [
    {
        "sheet_name": "Ablação_IA_Grids",
        "csv_path": os.path.join(BASE_DIR, "tabela_de_ablacao.csv"),
        "title": "Tabela de Ablação: Resolução de Grade Espacial vs. Algoritmos de Machine Learning",
        "script": "Untitled0.ipynb (Célula 13) / run_statistics.py",
        "descricao": "Avalia o impacto do tamanho da grade (1x1, 2x2, 3x3, 4x4) e algoritmos (Random Forest, Extra Trees, Gradient Boosting) nas métricas de teste cego (GroupShuffleSplit por animal)."
    },
    {
        "sheet_name": "Concordancia_Bootstrap",
        "csv_path": os.path.join(RD_DIR, "concordancia_bootstrap_cluster.csv"),
        "title": "Concordância Metrológica com Intervalos de Confiança por Bootstrap ao Nível de Animal",
        "script": "analise_estatistica_complementar.py (Seção 1)",
        "descricao": "Pearson r, CCC de Lin, MAE, RMSE, Viés, Beta OLS com IC95% empíricos (5.000 reamostragens por cluster de animal) e LoA de Bland-Altman para medidas repetidas."
    },
    {
        "sheet_name": "DeltaE_Descritivo",
        "csv_path": os.path.join(RD_DIR, "deltaE_descritivo.csv"),
        "title": "Distribuição de Delta E (CIE76) e Limiares de Equivalência por Tecido e Campanha",
        "script": "analise_estatistica_complementar.py (Seção 2)",
        "descricao": "Média, desvio padrão, mediana, quartis, P95, e percentual de leituras abaixo de 1.0 (imperceptível), 3.0 (equivalência laboratorial) e 5.0 (tolerância industrial)."
    },
    {
        "sheet_name": "Estabilidade_Animais",
        "csv_path": os.path.join(RD_DIR, "estabilidade_fatores_animais.csv"),
        "title": "Estabilidade da Concordância Frente a Fatores Biométricos do Animal (Idade, Peso, GMD)",
        "script": "analise_estatistica_complementar.py / script biométrico",
        "descricao": "Correlações de Pearson e Spearman entre os erros de predição e o avanço da idade (meses), peso corporal (kg) e ganho médio diário dos animais."
    },
    {
        "sheet_name": "Robustez_Ambiental",
        "csv_path": os.path.join(RD_DIR, "robustez_ambiental.csv"),
        "title": "Robustez às Condições Microclimáticas de Campo (TMR, CTR, Temp, UR, ITGU)",
        "script": "analise_estatistica_artigo.py (Seção 5)",
        "descricao": "Correlações de Pearson e Spearman entre erros absolutos / Delta E e as 5 variáveis meteorológicas registradas durante as coletas no curral."
    },
    {
        "sheet_name": "LMM_Periodo_LRT",
        "csv_path": os.path.join(RD_DIR, "lmm_efeito_periodo_LRT.csv"),
        "title": "Teste Omnibus da Razão de Verossimilhança (LRT) para Efeito Sazonal de Período",
        "script": "analise_estatistica_complementar.py (Seção 3)",
        "descricao": "Teste qui-quadrado de razão de verossimilhança (ML) do efeito fixo de campanha sazonal em modelos mistos com ajuste de p-valor por Holm."
    },
    {
        "sheet_name": "LMM_Vies_logDeltaE_Fixos",
        "csv_path": os.path.join(RD_DIR, "lmm_viés_logDeltaE_efeitos_fixos.csv"),
        "title": "Efeitos Fixos do Modelo Linear Misto: Viés com Sinal e log(Delta E)",
        "script": "analise_estatistica_complementar.py (Seção 3)",
        "descricao": "Estimativas fixas, erros-padrão, z, p-valores brutos e ajustados por Holm, IC 95% de Wald e ICC (Dezembro como período base)."
    },
    {
        "sheet_name": "LMM_Vies_logDeltaE_Diag",
        "csv_path": os.path.join(RD_DIR, "lmm_viés_logDeltaE_diagnostico.csv"),
        "title": "Diagnóstico de Resíduos do LMM (Viés com Sinal e log(Delta E))",
        "script": "analise_estatistica_complementar.py (Seção 3)",
        "descricao": "Teste de normalidade de Shapiro-Wilk (W e p), assimetria e curtose dos resíduos para comprovação das premissas gaussianas."
    },
    {
        "sheet_name": "Metricas_Classicas",
        "csv_path": os.path.join(RD_DIR, "metricas_concordancia.csv"),
        "title": "Métricas Clássicas Univariadas por Canal (App vs. Colorímetro)",
        "script": "analise_estatistica_artigo.py (Seção 2)",
        "descricao": "Pearson r, MAE, RMSE e Viés médio clássicos calculados por coordenada de cor sem modelagem de medidas repetidas."
    },
    {
        "sheet_name": "Regressao_Linear_Vies",
        "csv_path": os.path.join(RD_DIR, "regressao_linear.csv"),
        "title": "Regressão Linear OLS e Investigação de Viés Proporcional",
        "script": "analise_estatistica_artigo.py (Seção 3)",
        "descricao": "Ajuste y_app = alpha + beta * y_ref com testes de hipótese para alpha != 0 e beta != 1, SE do coeficiente e R²."
    },
    {
        "sheet_name": "Bland_Altman_Classico",
        "csv_path": os.path.join(RD_DIR, "bland_altman.csv"),
        "title": "Análise de Bland-Altman Clássica (Viés Médio e LoA 95% Ingênuos)",
        "script": "analise_estatistica_artigo.py (Seção 4)",
        "descricao": "Viés médio, DP das diferenças, limites inferior e superior de concordância ingênuos e correlação de proporcionalidade."
    },
    {
        "sheet_name": "Resumo_Geral_Integrado",
        "csv_path": os.path.join(RD_DIR, "resumo_geral.csv"),
        "title": "Tabela Resumo Geral Consolidando Métricas Clássicas",
        "script": "analise_estatistica_artigo.py (Seção 7)",
        "descricao": "Merge de concordância, regressão linear e Bland-Altman em visão única consolidada."
    },
    {
        "sheet_name": "LMM_ErroAbs_Resumo",
        "csv_path": os.path.join(RD_DIR, "modelo_linear_misto.csv"),
        "title": "Resumo de Ajuste do Modelo Misto sobre o Erro Absoluto",
        "script": "analise_estatistica_artigo.py (Seção 6)",
        "descricao": "N_obs, N_grupos, Log-Likelihood REML, Variância de Animal e Variância Residual para o modelo básico e modelo com covariáveis ambientais."
    },
    {
        "sheet_name": "LMM_ErroAbs_Efeitos_Fixos",
        "csv_path": os.path.join(RD_DIR, "modelo_linear_misto_efeitos_fixos.csv"),
        "title": "Efeitos Fixos Detalhados sobre o Erro Absoluto (com Correção Holm e BH-FDR)",
        "script": "analise_estatistica_artigo.py (Seção 6 atualizada)",
        "descricao": "Estimativas dos parâmetros para período e covariáveis ambientais z-score, com p-valores brutos, Holm e Benjamini-Hochberg."
    },
    {
        "sheet_name": "LMM_ErroAbs_Diagnostico",
        "csv_path": os.path.join(RD_DIR, "modelo_linear_misto_diagnostico.csv"),
        "title": "Diagnóstico de Resíduos do LMM sobre Erro Absoluto",
        "script": "analise_estatistica_artigo.py (Seção 6 atualizada)",
        "descricao": "Estatísticas de Shapiro-Wilk, assimetria e curtose demonstrando o desvio de normalidade nos erros absolutos brutos."
    }
]

# Estilos visuais openpyxl
font_title = Font(name="Calibri", size=14, bold=True, color="1F497D")
font_subtitle = Font(name="Calibri", size=10, italic=True, color="595959")
font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
font_data = Font(name="Calibri", size=10)
fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
fill_index_header = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
fill_zebra = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
border_thin = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

wb = openpyxl.Workbook()
# Remover sheet default
wb.remove(wb.active)

# -------------------------------------------------------------
# 1. ABA DE ÍNDICE E MAPEAMENTO DE SCRIPTS
# -------------------------------------------------------------
ws_index = wb.create_sheet(title="Índice_Metricas_Scripts")
ws_index.views.sheetView[0].showGridLines = True

ws_index.cell(row=1, column=1, value="CONSOLIDAÇÃO GERAL DE MÉTRICAS ESTATÍSTICAS – EXPERIMENTOS COLORÍMETRO (2025–2026)").font = font_title
ws_index.cell(row=2, column=1, value="Mapeamento de rastreabilidade entre abas, arquivos CSV de saída e scripts Python/Jupyter geradores").font = font_subtitle

index_headers = [
    "Item", "Nome da Aba", "Arquivo CSV de Origem", "Script / Código Gerador", "Descrição Científica das Métricas Contidas"
]

for col_idx, h in enumerate(index_headers, start=1):
    c = ws_index.cell(row=4, column=col_idx, value=h)
    c.font = font_header
    c.fill = fill_index_header
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for row_idx, meta in enumerate(TABLES_METADATA, start=5):
    ws_index.cell(row=row_idx, column=1, value=row_idx - 4).alignment = Alignment(horizontal="center")
    c_aba = ws_index.cell(row=row_idx, column=2, value=meta["sheet_name"])
    c_aba.font = Font(name="Calibri", size=10, bold=True, color="004B96")
    ws_index.cell(row=row_idx, column=3, value=os.path.basename(meta["csv_path"]))
    c_script = ws_index.cell(row=row_idx, column=4, value=meta["script"])
    c_script.font = Font(name="Calibri", size=10, bold=True, color="800000")
    ws_index.cell(row=row_idx, column=5, value=meta["descricao"])

    # Formatação e bordas
    is_even = (row_idx % 2 == 0)
    for c_idx in range(1, 6):
        cell = ws_index.cell(row=row_idx, column=c_idx)
        cell.border = border_thin
        if is_even:
            cell.fill = fill_zebra

# Ajuste de largura das colunas do índice
ws_index.column_dimensions['A'].width = 8
ws_index.column_dimensions['B'].width = 28
ws_index.column_dimensions['C'].width = 36
ws_index.column_dimensions['D'].width = 42
ws_index.column_dimensions['E'].width = 75
ws_index.row_dimensions[4].height = 28

# -------------------------------------------------------------
# 2. ABAS COM CADA CONJUNTO DE MÉTRICAS
# -------------------------------------------------------------
for meta in TABLES_METADATA:
    sheet_title = meta["sheet_name"]
    csv_file = meta["csv_path"]
    
    if not os.path.exists(csv_file):
        print(f"Aviso: {csv_file} não encontrado. Pulando...")
        continue
        
    df = pd.read_csv(csv_file)
    ws = wb.create_sheet(title=sheet_title)
    ws.views.sheetView[0].showGridLines = True
    
    # Cabeçalho informativo
    ws.cell(row=1, column=1, value=meta["title"]).font = font_title
    ws.cell(row=2, column=1, value=f"Gerado por: {meta['script']} | Arquivo original: {os.path.basename(csv_file)}").font = font_subtitle
    ws.cell(row=3, column=1, value=f"Escopo: {meta['descricao']}").font = font_subtitle
    
    # Headers dos dados
    header_row = 5
    ws.row_dimensions[header_row].height = 25
    for c_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=header_row, column=c_idx, value=col_name)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_thin

    # Linhas de dados
    for r_idx, row_values in enumerate(df.values, start=header_row + 1):
        is_even = (r_idx % 2 == 0)
        for c_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=r_idx, column=c_idx)
            
            # Tratar números e formatação
            if pd.isna(val):
                cell.value = "--"
            elif isinstance(val, (float, np.floating)):
                cell.value = float(val)
                if abs(val) < 0.0001 and val != 0:
                    cell.number_format = "0.00E+00"
                else:
                    cell.number_format = "#,##0.000"
            elif isinstance(val, (int, np.integer)):
                cell.value = int(val)
                cell.number_format = "#,##0"
            else:
                cell.value = str(val)
                
            cell.font = font_data
            cell.border = border_thin
            if is_even:
                cell.fill = fill_zebra
            cell.alignment = Alignment(horizontal="right" if isinstance(val, (int, float, np.number)) else "left", vertical="center")

    # Ajuste automático de largura de coluna
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < 5:  # Pular títulos no cálculo de largura
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

# Salvar planilha consolidada
wb.save(OUTPUT_XLSX)
print(f"\nPlanilha consolidada gerada com sucesso em:\n -> {OUTPUT_XLSX}")
print(f"Total de abas criadas: {len(wb.sheetnames)} (1 Índice + {len(wb.sheetnames)-1} tabelas de dados)")
